"""花名册照片导入（三种方式）+ 人脸入库。

方式一：向导——从幼师口袋同步名单后，逐个孩子拖照片。
方式二：Excel——两列：姓名 | 照片（照片列可以是文件路径，也可以是嵌入单元格的图片）。
方式三：整包照片自动匹配——把一堆照片丢进来，先用已有的库识别，认出来的自动归档，认不出的让老师指认。
"""
import hashlib
import io
import logging
import shutil
import time
import zipfile
from pathlib import Path
from typing import Optional

from PIL import Image, ImageOps

import db
import face_client
from config import PHOTO_DIR

log = logging.getLogger("roster")
IMG_EXT = {".jpg", ".jpeg", ".png", ".bmp", ".webp", ".heic"}


def pending_root() -> Path:
    """Unassigned images belong to the signed-in teacher, not a shared global tray."""
    uid = str(db.get_setting('current_uid') or '')
    if not uid.isdigit():
        raise ValueError('请先登录幼师口袋。')
    return PHOTO_DIR / '_pending' / uid


def _sha1(data: bytes) -> str:
    return hashlib.sha1(data).hexdigest()


def normalize_image(data: bytes, max_side: int = 1280) -> bytes:
    """统一转成 JPEG、限制边长（腾讯限 5MB，且小图更快）。"""
    im = Image.open(io.BytesIO(data))
    try:
        im = ImageOps.exif_transpose(im)
    except Exception:
        pass
    im = im.convert("RGB")
    if max(im.size) > max_side:
        im.thumbnail((max_side, max_side))
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=90)
    return buf.getvalue()


def save_student_photo(student_id: int, data: bytes, grade: str = None, score: float = None) -> dict:
    """保存到本地相册并登记；返回 photos 行。已存在（同哈希）则直接返回。
       grade: front=完整正脸 / ok=勉强能用 / None=老师手工传的，不评级。"""
    data = normalize_image(data)
    h = _sha1(data)
    ex = db.one("SELECT * FROM photos WHERE sha1=?", (h,))
    if ex:
        if ex['student_id'] != student_id:
            raise ValueError('这张照片已属于其他幼儿，请核对后选择该幼儿的独立照片。')
        # 同一张照片之前就存过（比如上一轮没评级就存进来了），这次知道等级了就补上
        if grade and not ex["grade"]:
            db.exec("UPDATE photos SET grade=?, quality=? WHERE id=?", (grade, score, ex["id"]))
            return db.one("SELECT * FROM photos WHERE id=?", (ex["id"],))
        return ex
    d = PHOTO_DIR / str(student_id)
    d.mkdir(parents=True, exist_ok=True)
    p = d / f"{h[:12]}.jpg"
    p.write_bytes(data)
    pid = db.exec("INSERT INTO photos(student_id,path,sha1,grade,quality,created_at) VALUES(?,?,?,?,?,?)",
                  (student_id, str(p), h, grade, score, time.time()))
    return db.one("SELECT * FROM photos WHERE id=?", (pid,))


def enroll_photo(photo_id: int, class_id: int) -> dict:
    """把一张本地照片入云端人脸库。"""
    ph = db.one("SELECT p.*, s.name, s.class_id FROM photos p JOIN students s ON s.student_id=p.student_id WHERE p.id=?", (photo_id,))
    if not ph:
        raise ValueError("照片不存在")
    if ph["enrolled"]:
        return ph
    data = Path(ph["path"]).read_bytes()
    try:
        fid = face_client.add_face(class_id, ph["student_id"], ph["name"], data)
        db.exec("UPDATE photos SET face_id=?, enrolled=1, error=NULL WHERE id=?", (fid, photo_id))
        n = db.one("SELECT COUNT(*) c FROM photos WHERE student_id=? AND enrolled=1", (ph["student_id"],))["c"]
        db.exec("UPDATE students SET face_count=?, face_person_id=? WHERE student_id=?", (n, face_client.person_id(ph["student_id"]), ph["student_id"]))
    except face_client.FaceError as e:
        db.exec("UPDATE photos SET error=? WHERE id=?", (str(e), photo_id))
        raise
    return db.one("SELECT * FROM photos WHERE id=?", (photo_id,))


def refresh_cloud_state(class_id: int) -> dict:
    """把云端的入库情况读回本地，让同班第二位老师直接看到绿点。

    只写 students.cloud_faces，不碰本地照片、不碰 enrolled —— 本地有几张还是几张，
    云端有几张单独记一笔。花名册上的点按"两边取大"来亮。
    """
    import time as _t
    counts = face_client.group_face_counts(class_id)
    sts = db.students_of(class_id)
    now = _t.time()
    hit = 0
    for s in sts:
        n = int(counts.get(s["student_id"], 0))
        if n:
            hit += 1
        db.exec("UPDATE students SET cloud_faces=?, cloud_at=? WHERE student_id=?",
                (n, now, s["student_id"]))
    return {"total": len(sts), "in_cloud": hit,
            "faces": sum(counts.values()), "at": now}


def enroll_pending(class_id: int, progress=None) -> dict:
    """把该班所有未入库照片入库。返回统计。"""
    cls = db.one("SELECT * FROM classes WHERE class_id=?", (class_id,))
    face_client.ensure_group(class_id, cls["name"] if cls else str(class_id))
    pend = db.rows("SELECT p.id FROM photos p JOIN students s ON s.student_id=p.student_id WHERE s.class_id=? AND p.enrolled=0", (class_id,))
    ok, fail, errors = 0, 0, []
    for i, r in enumerate(pend):
        try:
            enroll_photo(r["id"], class_id)
            ok += 1
        except Exception as e:
            fail += 1
            errors.append({"photo_id": r["id"], "error": str(e)})
        if progress:
            progress(i + 1, len(pend))
    return {"total": len(pend), "ok": ok, "fail": fail, "errors": errors}


# ---------- 方式二：Excel ----------
def import_excel(path: str, class_id: int) -> dict:
    """Excel 两列：姓名 | 照片。照片列支持：文件路径（绝对或相对 Excel 所在目录）、或嵌入图片。
       同一个孩子可以多行（多张照片）。返回 {matched:[{name,count}], unknown:[names], errors:[]}"""
    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    students = {s["name"]: s for s in db.students_of(class_id)}
    base = Path(path).parent
    # 表头识别
    header = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    name_col = next((i for i, h in enumerate(header) if "姓名" in h or "名字" in h or h.lower() == "name"), 0)
    photo_col = next((i for i, h in enumerate(header) if "照片" in h or "图片" in h or "相片" in h or h.lower() in ("photo", "image")), 1)
    # 嵌入图片：按行号归位
    embedded: dict[int, list[bytes]] = {}
    for img in getattr(ws, "_images", []):
        try:
            row = img.anchor._from.row + 1
            data = img._data()
            embedded.setdefault(row, []).append(data)
        except Exception as e:
            log.warning("读取嵌入图片失败：%s", e)
    result = {"matched": {}, "unknown": [], "errors": []}
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, name_col + 1).value
        if name is None:
            continue
        name = str(name).strip()
        st = students.get(name)
        if not st:
            if name not in result["unknown"]:
                result["unknown"].append(name)
            continue
        blobs: list[bytes] = list(embedded.get(r, []))
        cell = ws.cell(r, photo_col + 1).value
        if cell:
            for part in str(cell).replace("；", ";").replace("，", ",").replace(",", ";").split(";"):
                part = part.strip().strip('"')
                if not part:
                    continue
                p = Path(part)
                if not p.is_absolute():
                    p = base / p
                if p.is_dir():
                    for f in sorted(p.iterdir()):
                        if f.suffix.lower() in IMG_EXT:
                            blobs.append(f.read_bytes())
                elif p.exists():
                    blobs.append(p.read_bytes())
                else:
                    result["errors"].append(f"第{r}行 {name}：找不到照片 {part}")
        for b in blobs:
            try:
                save_student_photo(st["student_id"], b)
                result["matched"][name] = result["matched"].get(name, 0) + 1
            except Exception as e:
                result["errors"].append(f"第{r}行 {name}：{e}")
    result["matched"] = [{"name": k, "count": v} for k, v in result["matched"].items()]
    return result


# ---------- 方式三：文件夹 / 一堆照片自动归档 ----------
def import_folder_or_files(paths: list[str], class_id: int) -> dict:
    """规则：
       1) 若照片所在子文件夹名 == 孩子姓名，或文件名以孩子姓名开头（如 示例姓名1.jpg），直接归档；
       2) 否则若人脸库已有该班资料，用云端搜脸自动归档（得分≥阈值）；
       3) 剩下的进入"待指认"列表，由老师在界面上点名。"""
    students = db.students_of(class_id)
    by_name = {s["name"]: s for s in students}
    has_lib = any(s["enrolled_count"] for s in students)
    files: list[Path] = []
    for p in paths:
        p = Path(p)
        if p.is_dir():
            files += [f for f in p.rglob("*") if f.suffix.lower() in IMG_EXT]
        elif p.suffix.lower() == ".zip":
            tmp = PHOTO_DIR / "_unzip" / f"{int(time.time())}"
            tmp.mkdir(parents=True, exist_ok=True)
            with zipfile.ZipFile(p) as z:
                z.extractall(tmp)
            files += [f for f in tmp.rglob("*") if f.suffix.lower() in IMG_EXT]
        elif p.suffix.lower() in IMG_EXT:
            files.append(p)
    return import_blobs([(f.name, f.parent.name, f.read_bytes()) for f in files], class_id)


def import_blobs(blobs: list[tuple[str, str, bytes]], class_id: int) -> dict:
    """blobs: [(文件名, 所在文件夹名, 字节)]。逻辑同上。"""
    students = db.students_of(class_id)
    by_name = {s["name"]: s for s in students}
    has_lib = any(s["enrolled_count"] for s in students)
    res = {"auto": [], "pending": [], "errors": [], "total": len(blobs)}
    for fname, folder, data in blobs:
        f = Path(fname)
        try:
            st = by_name.get(folder)
            if not st:
                for nm, s in by_name.items():
                    if f.stem.startswith(nm):
                        st = s
                        break
            how = "name"
            score = None
            if not st and has_lib:
                try:
                    hits = face_client.search(class_id, normalize_image(data))
                    if hits and hits[0]["matched"]:
                        st = next((s for s in students if s["student_id"] == hits[0]["student_id"]), None)
                        score = hits[0]["score"]
                        how = "face"
                except Exception as e:
                    log.warning("搜脸失败 %s: %s", f, e)
            if st:
                ph = save_student_photo(st["student_id"], data)
                res["auto"].append({"file": str(f), "name": st["name"], "student_id": st["student_id"], "how": how, "score": score, "photo_id": ph["id"]})
            else:
                # 暂存到待指认区
                tmp = pending_root()
                tmp.mkdir(parents=True, exist_ok=True)
                nd = normalize_image(data)
                tp = tmp / f"{_sha1(nd)[:12]}.jpg"
                tp.write_bytes(nd)
                res["pending"].append({"file": str(f), "temp": str(tp)})
        except Exception as e:
            res["errors"].append(f"{f.name}: {e}")
    return res


def assign_pending(temp_path: str, student_id: int) -> dict:
    data = Path(temp_path).read_bytes()
    ph = save_student_photo(student_id, data)
    try:
        Path(temp_path).unlink()
    except Exception:
        pass
    return ph


def _delete_cloud_face(student_id: int, face_id: str):
    """删云端这张脸。腾讯不允许把一个人员的最后一张脸删光——
       真删到最后一张时，就把整个人员删掉，免得库里留着一张已经不要了的脸。"""
    n = db.one("SELECT COUNT(*) c FROM photos WHERE student_id=? AND enrolled=1", (student_id,))["c"]
    if n <= 1:
        face_client.delete_person(student_id)
    else:
        face_client.delete_face(student_id, face_id)


def delete_photo(photo_id: int):
    ph = db.one("SELECT * FROM photos WHERE id=?", (photo_id,))
    if not ph:
        return
    if ph["enrolled"] and ph["face_id"]:
        try:
            _delete_cloud_face(ph["student_id"], ph["face_id"])
        except Exception as e:
            log.warning("云端删脸失败：%s", e)
    try:
        Path(ph["path"]).unlink()
    except Exception:
        pass
    db.exec("DELETE FROM photos WHERE id=?", (photo_id,))
    n = db.one("SELECT COUNT(*) c FROM photos WHERE student_id=? AND enrolled=1", (ph["student_id"],))["c"]
    db.exec("UPDATE students SET face_count=? WHERE student_id=?", (n, ph["student_id"]))


# ---------- 从幼师口袋历史记录自动建库 ----------
def _collect_single_child_images(node, out: list, keep_multi: bool = False):
    """在返回的 JSON 里找记录里标注的孩子和它带的图片。
       默认只取"只标注了一个孩子"的记录 —— 那张照片是谁，老师当年已经标好了。
       keep_multi=True 时连合影记录也一起返回（out 里是 (sids, hashes)），
       留给第二轮"从合影里抠脸让老师指认"用。
       结构可能随站点变化，所以这里做的是宽松的树遍历，不写死字段路径。"""
    if isinstance(node, list):
        for v in node:
            _collect_single_child_images(v, out, keep_multi)
        return
    if not isinstance(node, dict):
        return

    sids = node.get("studentIds")
    if not sids:
        for key in ("studentVos", "studentVosWithMedal", "students"):
            v = node.get(key)
            if isinstance(v, list) and v and isinstance(v[0], dict):
                sids = [s.get("id") or s.get("studentId") for s in v]
                break
    if isinstance(sids, list):
        sids = [s for s in sids if isinstance(s, int)]

    hashes = []
    for key in ("fileHashWithEntityVos", "fileList", "fileVos", "files"):
        v = node.get(key)
        if isinstance(v, list):
            for fvo in v:
                if isinstance(fvo, dict) and (fvo.get("fileType") or "IMAGE") == "IMAGE":
                    h = fvo.get("hashName") or fvo.get("hash")
                    if h:
                        hashes.append(h)
                elif isinstance(fvo, str):
                    hashes.append(fvo)
    v = node.get("hashList")
    if isinstance(v, list):
        hashes += [h for h in v if isinstance(h, str)]

    if sids and hashes:
        if len(sids) == 1:
            out.append((sids[0], hashes) if not keep_multi else (sids, hashes))
        elif keep_multi:
            out.append((sids, hashes))

    for v in node.values():
        _collect_single_child_images(v, out, keep_multi)


MAX_PROBE_PER_CHILD = 24      # 每个孩子最多验多少张候选（每验一张 = 1 次腾讯人脸检测）


def harvest_from_ysk(class_id: int, ysk, want: int = 2, max_pages: int = 40,
                     reset: bool = False, progress=None) -> dict:
    """从幼师口袋这个班的历史观察记录里，自动给每个孩子找底照。

    只采用「只标注了一个孩子」的记录 —— 那张照片是谁，老师当年已经标好了。
    下载后再用腾讯云做一次人脸检测，只留"画面里正好一张脸"的，避免把同框的
    其他孩子当成本人入库。返回统计给界面显示。
    """
    if reset:                       # 重挑：先把自动找来的底照清掉（老师手工传的保留）
        for ph in db.rows("""SELECT p.id FROM photos p JOIN students s ON s.student_id=p.student_id
                             WHERE s.class_id=? AND p.grade IS NOT NULL""", (class_id,)):
            try:
                delete_photo(ph["id"])
            except Exception as e:
                log.warning("清理旧底照失败 %s: %s", ph["id"], e)
    students = {s["student_id"]: s for s in db.students_of(class_id)}
    need = {sid: max(0, want - s["photo_count"]) for sid, s in students.items()}
    res = {"pages": 0, "records": 0, "candidates": 0, "saved": 0,
           "skipped_multiface": 0, "skipped_noface": 0, "skipped_sideface": 0, "failed": 0,
           "front_count": 0, "no_front": [],
           "per_child": {}, "still_missing": [], "errors": []}
    if not any(need.values()):
        res["still_missing"] = []
        return res

    # 1) 翻页收集候选（只读接口，不下载图片，很快）
    cand: dict[int, list[str]] = {}
    for page in range(1, max_pages + 1):
        try:
            data = ysk.class_records(class_id, page)
        except Exception as e:
            res["errors"].append(f"第 {page} 页读取失败：{e}")
            break
        pairs: list = []
        _collect_single_child_images(data, pairs)
        res["pages"] = page
        if not pairs:
            break
        res["records"] += len(pairs)
        for sid, hashes in pairs:
            if sid not in students:
                continue
            lst = cand.setdefault(sid, [])
            for h in hashes:
                if h not in lst:
                    lst.append(h)
        if progress:
            progress(page, max_pages)
        if all(len(cand.get(sid, [])) >= need[sid] + 8 for sid in need if need[sid] > 0):
            break   # 每个还缺照片的孩子都攒够候选了，不用再翻
    res["candidates"] = sum(len(v) for v in cand.values())

    # 2) 逐个孩子挑底照：只要"完整正脸"，这张不合格就换下一张，一直往下翻。
    #    实在一张完整正脸都找不到，才退而用"勉强能用"的，并在结果里点名让老师换。
    todo = [sid for sid in cand if need.get(sid, 0) > 0]
    for i, sid in enumerate(todo):
        name = students[sid]["name"]
        fronts, oks = [], []
        for h in cand[sid][:MAX_PROBE_PER_CHILD]:
            if len(fronts) >= need[sid]:
                break                      # 已经凑够完整正脸，不用再翻了
            try:
                data = normalize_image(ysk.download(ysk.image_url(h)))
            except Exception as e:
                res["failed"] += 1
                res["errors"].append(f"{name}：下载失败 {e}")
                continue
            try:
                faces = face_client.detect_faces(data)
            except Exception as e:
                msg = str(e)
                if "没有检测到人脸" in msg or "NoFaceInPhoto" in msg:
                    res["skipped_noface"] += 1
                elif "FaceSizeTooSmall" in msg or "人脸太小" in msg:
                    res["skipped_sideface"] += 1      # 脸太小，同样是"这张不能用"
                else:
                    res["errors"].append(f"{name}：人脸检测失败 {e}")
                continue
            if not faces:
                res["skipped_noface"] += 1
                continue
            if len(faces) > 1:
                res["skipped_multiface"] += 1      # 合影不敢用，不知道哪个是他
                continue
            f = faces[0]
            g = face_client.face_grade(f)
            sc = face_client.frontal_score(f)
            if g == "front":
                fronts.append((sc, data))
            elif g == "ok":
                oks.append((sc, data, face_client.why_not_front(f)))
                res["skipped_sideface"] += 1
            else:
                res["skipped_sideface"] += 1
        fronts.sort(key=lambda x: -x[0])
        oks.sort(key=lambda x: -x[0])

        picked = [(sc, d, "front", None) for sc, d in fronts[:need[sid]]]
        if len(picked) < need[sid]:            # 正脸不够，用勉强能用的补位
            for sc, d, why in oks[:need[sid] - len(picked)]:
                picked.append((sc, d, "ok", why))
        got = 0
        for sc, data, g, why in picked:
            try:
                save_student_photo(sid, data, grade=g, score=round(sc, 1))
                got += 1
                res["saved"] += 1
                if g == "front":
                    res["front_count"] += 1
            except Exception as e:
                res["failed"] += 1
                res["errors"].append(f"{name}：保存失败 {e}")
        if got:
            res["per_child"][name] = got
        if not fronts and got:
            res["no_front"].append(f"{name}（{picked[0][3] or '不够正'}）")
        if progress:
            progress(max_pages + i + 1, max_pages + len(todo))

    for sid, s in students.items():
        n = db.one("SELECT COUNT(*) c FROM photos WHERE student_id=?", (sid,))["c"]
        if n == 0:
            res["still_missing"].append(s["name"])
    res["errors"] = res["errors"][:15]
    return res


def regrade_photos(class_id: int, progress=None) -> dict:
    """给还没评过级的底照补一次"正脸判定"，让复核页面能显示 ✓正脸 / ⚠不够正。"""
    ph = db.rows("""SELECT p.id, p.path, s.name FROM photos p JOIN students s ON s.student_id=p.student_id
                    WHERE s.class_id=? AND p.grade IS NULL""", (class_id,))
    res = {"total": len(ph), "front": 0, "ok": 0, "bad": 0, "failed": 0, "bad_list": [], "errors": []}
    for i, r in enumerate(ph):
        try:
            faces = face_client.detect_faces(Path(r["path"]).read_bytes())
        except Exception as e:
            res["failed"] += 1
            res["errors"].append(f"{r['name']}：{e}")
            continue
        if not faces:
            g, sc = "bad", 0.0
        else:
            f = max(faces, key=face_client.frontal_score)
            g, sc = face_client.face_grade(f), face_client.frontal_score(f)
            if len(faces) > 1:
                g = "bad"                     # 合影，不该当底照
        db.exec("UPDATE photos SET grade=?, quality=? WHERE id=?", (g, round(sc, 1), r["id"]))
        res[g] = res.get(g, 0) + 1
        if g == "bad":
            res["bad_list"].append(r["name"])
        if progress:
            progress(i + 1, len(ph))
    res["errors"] = res["errors"][:15]
    return res


def _crop_face(data: bytes, f: dict, margin: float = 0.55) -> bytes:
    """从合影里把一张脸连同周围一圈裁出来（留白多一点，人脸库需要额头下巴都在）。"""
    im = Image.open(io.BytesIO(data)).convert("RGB")
    W, H = im.size
    x, y, w, h = f.get("x", 0), f.get("y", 0), f.get("w", 0), f.get("h", 0)
    mx, my = int(w * margin), int(h * margin)
    box = (max(0, x - mx), max(0, y - int(my * 1.2)), min(W, x + w + mx), min(H, y + h + my))
    im = im.crop(box)
    if max(im.size) < 300:                     # 裁出来太小的话放大一点，云端才认
        r = 300 / max(im.size)
        im = im.resize((int(im.width * r), int(im.height * r)), Image.LANCZOS)
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=92)
    return buf.getvalue()


def save_candidate(student_id: int, data: bytes, grade: str, score: float, src: str = "group") -> Optional[dict]:
    """存一张"疑似是这个孩子"的待指认照片，等老师点头再入库。"""
    h = _sha1(data)
    if db.one("SELECT 1 FROM candidates WHERE sha1=?", (h,)) or db.one("SELECT 1 FROM photos WHERE sha1=?", (h,)):
        return None
    d = PHOTO_DIR / "_cand" / str(student_id)
    d.mkdir(parents=True, exist_ok=True)
    p = d / f"{h[:12]}.jpg"
    p.write_bytes(data)
    cid = db.exec("INSERT INTO candidates(student_id,path,sha1,score,grade,src,created_at) VALUES(?,?,?,?,?,?,?)",
                  (student_id, str(p), h, score, grade, src, time.time()))
    return db.one("SELECT * FROM candidates WHERE id=?", (cid,))


def confirm_candidate(cand_id: int) -> dict:
    """老师确认这张就是本人 → 转成正式底照。"""
    c = db.one("SELECT * FROM candidates WHERE id=?", (cand_id,))
    if not c:
        raise ValueError("候选照片不存在")
    ph = save_student_photo(c["student_id"], Path(c["path"]).read_bytes(), grade=c["grade"], score=c["score"])
    drop_candidate(cand_id)
    return ph


def drop_candidate(cand_id: int):
    c = db.one("SELECT * FROM candidates WHERE id=?", (cand_id,))
    if not c:
        return
    try:
        Path(c["path"]).unlink()
    except Exception:
        pass
    db.exec("DELETE FROM candidates WHERE id=?", (cand_id,))


def class_candidates(class_id: int) -> list[dict]:
    return db.rows("""SELECT c.id, c.student_id, c.score, c.grade, c.src, s.name
                      FROM candidates c JOIN students s ON s.student_id=c.student_id
                      WHERE s.class_id=? ORDER BY s.name, c.score DESC""", (class_id,))


MAX_GROUP_PROBE = 10          # 每个孩子最多翻多少张合影


def harvest_from_group_photos(class_id: int, ysk, max_pages: int = 60, progress=None) -> dict:
    """第二轮：还没底照（或还没一张正脸）的孩子，去合影里找。

    合影本身不知道哪张脸是谁，所以做法是"排除法 + 老师点头"：
      1) 把合影里每张脸单独抠出来；
      2) 用已经建好的人脸库搜一遍 —— 认出来是别的孩子的，直接排除；
         认出来就是本人的，直接入库；
      3) 剩下认不出的，放进"待指认"，界面上让老师点一下确认。
    """
    students = {s["student_id"]: s for s in db.students_of(class_id)}
    res = {"pages": 0, "targets": [], "auto": 0, "candidates": 0,
           "skipped_matched_other": 0, "skipped_bad": 0, "errors": []}

    # 谁需要：一张照片都没有的，或者一张"正脸"都没有的
    need = []
    for sid, s in students.items():
        n_front = db.one("SELECT COUNT(*) c FROM photos WHERE student_id=? AND (grade='front' OR grade IS NULL)", (sid,))["c"]
        if n_front == 0:
            need.append(sid)
    if not need:
        return res
    res["targets"] = [students[sid]["name"] for sid in need]

    # 收集合影候选
    cand: dict[int, list[str]] = {sid: [] for sid in need}
    for page in range(1, max_pages + 1):
        try:
            data = ysk.class_records(class_id, page)
        except Exception as e:
            res["errors"].append(f"第 {page} 页读取失败：{e}")
            break
        pairs: list = []
        _collect_single_child_images(data, pairs, keep_multi=True)
        res["pages"] = page
        if not pairs:
            break
        for sids, hashes in pairs:
            if isinstance(sids, int):
                sids = [sids]
            for sid in sids:
                if sid in cand:
                    for h in hashes:
                        if h not in cand[sid]:
                            cand[sid].append(h)
        if progress:
            progress(page, max_pages)
        if all(len(cand[sid]) >= MAX_GROUP_PROBE for sid in need):
            break

    done = set()
    for i, sid in enumerate(need):
        name = students[sid]["name"]
        got = 0
        for h in cand[sid][:MAX_GROUP_PROBE]:
            if got >= 2:
                break
            try:
                raw = normalize_image(ysk.download(ysk.image_url(h, 1600)), max_side=1600)
                faces = face_client.detect_faces(raw)
            except Exception as e:
                res["errors"].append(f"{name}：{e}")
                continue
            for f in faces:
                if got >= 2:
                    break
                if face_client.face_grade(f) == "bad":
                    res["skipped_bad"] += 1
                    continue
                try:
                    crop = _crop_face(raw, f)
                    hits = face_client.search(class_id, crop, max_faces=1)
                except Exception as e:
                    res["errors"].append(f"{name}：抠脸/搜脸失败 {e}")
                    continue
                sc = face_client.frontal_score(f)
                g = face_client.face_grade(f)
                best = hits[0] if hits else None
                if best and best.get("matched"):
                    if best["student_id"] == sid:
                        try:            # 库里已经认得他，直接补一张
                            save_student_photo(sid, crop, grade=g, score=round(sc, 1))
                            res["auto"] += 1
                            got += 1
                        except Exception as e:
                            res["errors"].append(f"{name}：保存失败 {e}")
                    else:
                        res["skipped_matched_other"] += 1     # 是同框的别的孩子
                    continue
                if save_candidate(sid, crop, g, round(sc, 1)):
                    res["candidates"] += 1
                    got += 1
        if got:
            done.add(sid)
        if progress:
            progress(max_pages + i + 1, max_pages + len(need))
    res["errors"] = res["errors"][:15]
    return res


def reassign_photo(photo_id: int, new_student_id: int) -> dict:
    """老师复核时发现这张照片认错人了，把它改派给正确的孩子。
       会先把云端那张脸删掉（如果已入库），改派后重新入库即可。"""
    ph = db.one("SELECT * FROM photos WHERE id=?", (photo_id,))
    if not ph:
        raise ValueError("照片不存在")
    if ph["student_id"] == new_student_id:
        return ph
    if ph["enrolled"] and ph["face_id"]:
        try:
            _delete_cloud_face(ph["student_id"], ph["face_id"])
        except Exception as e:
            log.warning("云端删脸失败：%s", e)
    old = ph["student_id"]
    d = PHOTO_DIR / str(new_student_id)
    d.mkdir(parents=True, exist_ok=True)
    newp = d / Path(ph["path"]).name
    try:
        shutil.move(ph["path"], newp)
    except Exception:
        newp = Path(ph["path"])          # 移动失败就原地放着，不影响使用
    db.exec("UPDATE photos SET student_id=?, path=?, enrolled=0, face_id=NULL, error=NULL WHERE id=?",
            (new_student_id, str(newp), photo_id))
    for sid in (old, new_student_id):
        n = db.one("SELECT COUNT(*) c FROM photos WHERE student_id=? AND enrolled=1", (sid,))["c"]
        db.exec("UPDATE students SET face_count=? WHERE student_id=?", (n, sid))
    return db.one("SELECT * FROM photos WHERE id=?", (photo_id,))


def class_photos(class_id: int) -> list[dict]:
    """整班的底照，供"复核底照"页用。"""
    return db.rows("""SELECT p.id, p.student_id, p.enrolled, p.error, p.grade, p.quality, s.name
                      FROM photos p JOIN students s ON s.student_id=p.student_id
                      WHERE s.class_id=? ORDER BY s.name, p.id""", (class_id,))
